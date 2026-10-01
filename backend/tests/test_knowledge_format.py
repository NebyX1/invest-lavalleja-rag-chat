import json

import pytest
from docx import Document

from ingest import build_chunks
from knowledge_format import SCHEMA, export_docx, read_jsonl


def test_docx_conversion_preserves_tables_metadata_and_excludes_internal_appendix(tmp_path):
    doc = Document()
    doc.add_paragraph("Z1 / MINAS")
    doc.add_heading("O01 · Servicios logísticos", level=1)
    doc.add_paragraph("Conocimiento público " * 30)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Contacto"
    table.cell(0, 1).text = "Correo"
    table.cell(1, 0).text = "Oficina publicada"
    table.cell(1, 1).text = "publico@example.test"
    doc.add_paragraph("ANEXO INTERNO / RAG")
    doc.add_paragraph("SECRETO QUE NO DEBE INDEXARSE")
    source, target = tmp_path / "guia.docx", tmp_path / "guia.jsonl"
    doc.save(source)
    export_docx(source, target)
    manifest, chunks = read_jsonl(target)
    assert manifest["schema"] == SCHEMA
    assert chunks == build_chunks(source)
    assert "SECRETO" not in target.read_text(encoding="utf-8")
    assert any("Correo: publico@example.test" in c["text"] for c in chunks)
    assert chunks[0]["zone"] == "Z1" and chunks[0]["card"] == "O01"


@pytest.mark.parametrize("bad", [
    {"section": "ANEXO INTERNO > Privado"}, {"zone": "Z1' OR 1=1"},
    {"text": ""}, {"doc": "../../secret"}, {"card": "otro"},
])
def test_jsonl_rejects_internal_content_or_invalid_metadata(tmp_path, bad):
    chunk = {"type": "chunk", "section": "Público", "text": "Texto", "zone": "", "card": "", "doc": "guia", **bad}
    target = tmp_path / "bad.jsonl"
    target.write_text(json.dumps({"type": "manifest", "schema": SCHEMA, "chunk_count": 1}) + "\n" + json.dumps(chunk), encoding="utf-8")
    with pytest.raises(ValueError):
        read_jsonl(target)
